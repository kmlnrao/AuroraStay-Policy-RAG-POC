import csv
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


TEMPLATE = Path("attached_assets/Presentation_Template_1789946721847.pptx")
OUTPUT = Path("Project_3_Delivery_Exception_Resolution_Presentation.pptx")
DETAIL_CSV = Path("delivery_exception_evaluation_detail.csv")

prs = Presentation(TEMPLATE)

NAVY = RGBColor(24, 55, 87)
BLUE = RGBColor(45, 111, 166)
TEAL = RGBColor(0, 139, 139)
GREEN = RGBColor(35, 130, 84)
ORANGE = RGBColor(210, 115, 35)
GRAY = RGBColor(82, 95, 108)
LIGHT = RGBColor(242, 247, 250)
WHITE = RGBColor(255, 255, 255)


def remove_all(slide):
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.25), Inches(8.9), Inches(0.55))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Aptos Display"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.58), Inches(0.78), Inches(8.6), Inches(0.3))
        stf = sb.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        sp.text = subtitle
        sp.font.name = "Aptos"
        sp.font.size = Pt(9)
        sp.font.color.rgb = GRAY


def add_footer(slide, text="Project 3 | Low-Code Delivery Exception Resolution"):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(10), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.08), Inches(8.9), Inches(0.2))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(7)
    p.font.color.rgb = GRAY


def add_body(slide, text, x=0.7, y=1.25, w=8.65, h=5.4, size=15, color=GRAY,
             bullet=True, spacing=5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.05)
    lines = text if isinstance(text, list) else text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.level = 0
        p.font.bold = line.endswith(":") or line.startswith("RESULT")
        if bullet and line.strip() and not line.startswith("RESULT"):
            p.text = "• " + line
    return box


def add_card(slide, title, body, x, y, w, h, accent=TEAL, size=11):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT
    card.line.color.rgb = RGBColor(210, 224, 232)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    band.fill.solid()
    band.fill.fore_color.rgb = accent
    band.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.12), Inches(w - 0.28), Inches(h - 0.2))
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Aptos"
    p.font.size = Pt(size + 1)
    p.font.bold = True
    p.font.color.rgb = NAVY
    for line in body.split("\n"):
        q = tf.add_paragraph()
        q.text = line
        q.font.name = "Aptos"
        q.font.size = Pt(size)
        q.font.color.rgb = GRAY
        q.space_after = Pt(2)
    return card


def replace_text(shape, text, size=18, color=NAVY, bold=False):
    shape.text = text
    for p in shape.text_frame.paragraphs:
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold


# Slide 1 - title
slide = prs.slides[0]
remove_all(slide)
bg = slide.background.fill
bg.solid()
bg.fore_color.rgb = NAVY
accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.28), Inches(7.5))
accent.fill.solid()
accent.fill.fore_color.rgb = TEAL
accent.line.fill.background()
tb = slide.shapes.add_textbox(Inches(0.9), Inches(1.25), Inches(8.2), Inches(1.2))
tf = tb.text_frame
p = tf.paragraphs[0]
p.text = "AI-Powered Delivery\nException Resolution"
p.font.name = "Aptos Display"
p.font.size = Pt(32)
p.font.bold = True
p.font.color.rgb = WHITE
sub = slide.shapes.add_textbox(Inches(0.95), Inches(2.85), Inches(7.5), Inches(0.7))
sp = sub.text_frame.paragraphs[0]
sp.text = "Project 3 | Low-Code Multi-Agent System"
sp.font.name = "Aptos"
sp.font.size = Pt(19)
sp.font.color.rgb = RGBColor(196, 224, 232)
date = slide.shapes.add_textbox(Inches(0.95), Inches(5.95), Inches(5.5), Inches(0.45))
dp = date.text_frame.paragraphs[0]
dp.text = "20 September 2026"
dp.font.name = "Aptos"
dp.font.size = Pt(13)
dp.font.color.rgb = WHITE

# Slide 2 - agenda
slide = prs.slides[1]
remove_all(slide)
add_title(slide, "Contents / Agenda")
agenda = [
    "Executive summary and verified outcomes",
    "Business problem and solution approach",
    "Data loading, model initialization, and retrieval",
    "Tools, multi-agent architecture, escalation, and security",
    "Test cases, traces, document citations, and evaluation metrics",
    "Conclusions, recommendations, and appendix",
]
add_body(slide, agenda, x=0.9, y=1.45, w=8.2, h=4.8, size=18, color=NAVY, spacing=12)
add_footer(slide)

# Slide 3 - executive summary
slide = prs.slides[2]
remove_all(slide)
add_title(slide, "Executive Summary", "The completed notebook was executed against all attached data.")
add_card(slide, "SCOPE", "13 raw delivery-log rows\n10 unique shipments\n12 customers + 6 lockers\n10-page operations playbook", 0.65, 1.25, 2.75, 1.65, BLUE, 11)
add_card(slide, "RESULT", "Task completion: 100%\nEscalation accuracy: 100% (8/8)\nTool-call accuracy: 100%\nCoherence: 5.00 / 5.0", 3.62, 1.25, 2.75, 1.65, GREEN, 11)
add_card(slide, "RECOMMENDATION", "Keep deterministic policy rules authoritative for safety and escalation. Use the LLM for classification, rationale, and customer communication.", 6.59, 1.25, 2.75, 1.65, ORANGE, 11)
add_body(slide, [
    "The system filters routine noise before unnecessary tool or model calls.",
    "It removes duplicate scans, retrieves policy context, checks lockers, and applies escalation rules.",
    "Structured Resolution and Communication agents are validated by critic agents before finalization.",
    "The run used a credential-free deterministic fallback; remote GPT-4o-mini mode remains supported when configured.",
], x=0.85, y=3.35, w=8.3, h=2.8, size=15, color=GRAY, spacing=8)
add_footer(slide)

# Slide 4 - problem and approach
slide = prs.slides[3]
remove_all(slide)
add_title(slide, "Business Problem Overview and Solution Approach")
add_card(slide, "BUSINESS PROBLEM", "Delivery exceptions are messy, repetitive, and policy-sensitive. Duplicate scans create false work; natural-language notes hide the relevant facts; inconsistent escalation risks customer and operational harm.", 0.65, 1.25, 4.1, 2.1, ORANGE, 13)
add_card(slide, "SOLUTION", "A hub-and-spoke LangGraph workflow combines guardrails, deterministic tools, structured agent outputs, critic validation, and traceable final actions.", 5.05, 1.25, 4.25, 2.1, TEAL, 13)
add_body(slide, [
    "1. Preprocess: deduplicate, consolidate attempts, scan for injection, and filter routine noise.",
    "2. Enrich: retrieve customer profile, locker availability, playbook chunks, and escalation signals.",
    "3. Decide: Resolution Agent selects classification and action with a policy-grounded rationale.",
    "4. Validate: Resolution Critic may accept, revise, or escalate; Communication Critic validates the message.",
    "5. Finalize: package an auditable action, escalation flag, customer message, citations, and latency.",
], x=0.8, y=3.7, w=8.55, h=2.55, size=14, color=GRAY, spacing=5)
add_footer(slide)

# Slide 5 - data
slide = prs.slides[4]
remove_all(slide)
add_title(slide, "Data Loading", "Source coverage and attributes used by the solution")
add_card(slide, "DELIVERY LOGS | 13 ROWS", "shipment_id, timestamp, status_code, driver description, customer_id, address, package type/size, attempt number, duplicate flag", 0.55, 1.2, 4.3, 1.7, BLUE, 11)
add_card(slide, "CUSTOMER DB | 12 ROWS", "Customer tier, preferred channel, exception count in the last 90 days, active credit, and name. Names are redacted by default.", 5.05, 1.2, 4.3, 1.7, TEAL, 11)
add_card(slide, "LOCKER DB | 6 ROWS", "Locker address and ZIP, capacity status, operating hours, and maximum package size. Eligibility is checked before reroute communication.", 0.55, 3.15, 4.3, 1.7, GREEN, 11)
add_card(slide, "PLAYBOOK + LABELS", "10-page PDF split into 24 page-aware chunks. Ground truth has 13 labels aligned to the raw logs and consolidates to 10 unique shipment scenarios.", 5.05, 3.15, 4.3, 1.7, ORANGE, 11)
add_body(slide, [
    "Important data quality handling: duplicate scans for SHP-002 and SHP-007 are removed before consolidation.",
    "Ground truth is consolidated by shipment, using the highest-attempt exception label for multi-row shipments.",
], x=0.75, y=5.45, w=8.5, h=0.95, size=12, color=GRAY, spacing=3)
add_footer(slide)

# Slide 6 - models and runtime
slide = prs.slides[5]
remove_all(slide)
add_title(slide, "Model Initialization and Retrieval")
add_card(slide, "REMOTE MODE", "Generation: gpt-4o-mini\nValidation: gpt-4o-mini\nTemperature: 0.0\nCredentials loaded only from config.json or environment", 0.65, 1.25, 2.75, 2.0, BLUE, 12)
add_card(slide, "VALIDATION MODE", "The executed submission used a deterministic offline fallback because no credentials were attached. This makes the evaluation reproducible and still exercises the complete graph.", 3.62, 1.25, 2.75, 2.0, GREEN, 12)
add_card(slide, "RETRIEVAL", "Page-aware lexical retrieval over 24 chunks. It returns source page metadata for audit citations without requiring a model download.", 6.59, 1.25, 2.75, 2.0, TEAL, 12)
add_body(slide, [
    "Pydantic schemas enforce valid Resolution, Communication, and Critic outputs.",
    "Temperature 0.0 is used to reduce variance in structured decisions.",
    "The notebook supports remote LLM execution when a valid config.json is supplied, but never hard-codes credentials.",
    "The offline path is explicitly labeled in the exported results as runtime_mode=offline_fallback.",
], x=0.8, y=3.8, w=8.5, h=2.1, size=15, color=GRAY, spacing=7)
add_footer(slide)

# Slide 7 - tools
slide = prs.slides[6]
remove_all(slide)
add_title(slide, "Defining Tools")
tools = [
    ("Delivery Log Reader", "Reads all raw CSV rows and feeds the preprocessor."),
    ("Customer Profile Lookup", "Parameterized SQLite query; PII name is returned only with explicit include_pii=True."),
    ("Locker Availability Checker", "Checks ZIP, size capacity, FULL/LIMITED status, and returns eligibility plus reason."),
    ("Playbook Search", "Retrieves relevant page-aware policy chunks for the Resolution and Critic agents."),
    ("Escalation Rule Engine", "Deterministically evaluates automatic and discretionary triggers, including delay-hour parsing."),
]
for i, (title, body) in enumerate(tools):
    y = 1.15 + i * 0.95
    add_card(slide, title, body, 0.75, y, 8.45, 0.72, [BLUE, TEAL, GREEN, ORANGE, NAVY][i], 11)
add_body(slide, [
    "Evaluation result: 10/10 scenarios invoked the expected tool path, including the intended noise short-circuit.",
], x=0.8, y=6.15, w=8.5, h=0.4, size=12, color=GREEN, spacing=2)
add_footer(slide)

# Slide 8 - architecture
slide = prs.slides[7]
remove_all(slide)
add_title(slide, "Multi-Agent System Architecture")
add_body(slide, [
    "Hub-and-spoke orchestration keeps routing deterministic while allowing agent outputs to be structured and critic-checked.",
], x=0.75, y=1.0, w=8.5, h=0.35, size=12, color=GRAY, bullet=False, spacing=2)

nodes = [
    ("Preprocessor", 0.45, 2.05, BLUE),
    ("Orchestrator", 2.25, 2.05, NAVY),
    ("Resolution\nAgent", 4.05, 1.45, TEAL),
    ("Resolution\nCritic", 4.05, 2.7, ORANGE),
    ("Communication\nAgent", 6.0, 1.45, GREEN),
    ("Communication\nCritic", 6.0, 2.7, ORANGE),
    ("Finalize", 7.95, 2.05, NAVY),
]
for label, x, y, color in nodes:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.35), Inches(0.75))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = WHITE
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = label
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Aptos"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = WHITE
for x1, y1, x2, y2 in [
    (1.8, 2.42, 2.25, 2.42), (3.6, 2.42, 4.05, 1.82),
    (3.6, 2.42, 4.05, 3.07), (5.4, 1.82, 6.0, 1.82),
    (5.4, 3.07, 6.0, 3.07), (7.35, 1.82, 7.95, 2.42),
]:
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = GRAY
    conn.line.width = Pt(1.5)
add_body(slide, [
    "Prompts: policy-grounded resolution, tier-aware communication, resolution validation, communication validation, and 1–5 trajectory coherence judging.",
    "Escalation: automatic triggers (third attempt, VIP ≥3 exceptions, damaged perishables, suspected fraud, safety) plus discretionary triggers (high-history Standard and delayed Premium perishables).",
    "Security: parameterized SQL, default PII redaction, prompt-injection scanning before LLM calls, no raw PII in coherence judging, and safe fallback escalation.",
], x=0.8, y=4.2, w=8.5, h=2.2, size=12, color=GRAY, spacing=5)
add_footer(slide)

# Slide 9 - testing and aggregate metrics
slide = prs.slides[8]
remove_all(slide)
add_title(slide, "Test Cases Execution, Evaluation, and Observability")
add_card(slide, "AGGREGATE METRICS", "Task completion 100%\nEscalation accuracy 100% (8/8)\nTool-call accuracy 100%\nCoherence 5.00/5.0", 0.65, 1.15, 2.7, 1.75, GREEN, 12)
add_card(slide, "OBSERVABILITY", "Every exception trace records tool calls, agent decisions, critic decisions, escalation signals, latency, and retrieved playbook page citations.", 3.65, 1.15, 2.7, 1.75, BLUE, 12)
add_card(slide, "ACCURACY COMMENT", "All expected resolutions, tones, exception classifications, and escalation decisions match the supplied ground truth for the 10 unique scenarios.", 6.65, 1.15, 2.7, 1.75, TEAL, 12)
test_lines = [
    "SHP-001 Delivered → noise / no action",
    "SHP-002 VIP attempts → reschedule / escalate",
    "SHP-003 Address issue → reschedule",
    "SHP-004 VIP fragile damage → replace / escalate",
    "SHP-005 Third attempt → locker reroute / escalate",
    "SHP-006 Refused → return to sender",
    "SHP-007 Perishable weather >4h → replace / escalate",
    "SHP-008 Standard high history → reschedule / escalate",
    "SHP-009 Damaged perishable → replace / escalate",
    "SHP-010 Depot scan → noise / no action",
]
add_body(slide, test_lines, x=0.8, y=3.35, w=8.5, h=3.2, size=12, color=GRAY, spacing=2)
add_footer(slide)

# Slide 10 - appendix divider
slide = prs.slides[9]
remove_all(slide)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(5.85), Inches(10), Inches(1.1))
bar.fill.solid()
bar.fill.fore_color.rgb = TEAL
bar.line.fill.background()
tb = slide.shapes.add_textbox(Inches(0.75), Inches(2.2), Inches(8.5), Inches(1))
p = tb.text_frame.paragraphs[0]
p.text = "APPENDIX"
p.font.name = "Aptos Display"
p.font.size = Pt(34)
p.font.bold = True
p.font.color.rgb = WHITE
sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.25), Inches(8.2), Inches(0.6))
sp = sub.text_frame.paragraphs[0]
sp.text = "Detailed scenario outcomes, citations, and submission checklist"
sp.font.name = "Aptos"
sp.font.size = Pt(16)
sp.font.color.rgb = RGBColor(196, 224, 232)

# Slide 11 - detailed results table
slide = prs.slides[10]
remove_all(slide)
add_title(slide, "Appendix: Detailed Evaluation Results")
headers = ["Shipment", "Decision", "Esc.", "Tone", "Task", "Citations"]
rows = []
with DETAIL_CSV.open(newline="") as f:
    for row in csv.DictReader(f):
        decision = row["Resolution (Ground Truth/Prediction)"].split("/")[-1]
        esc = row["Escalation (Ground Truth/Prediction)"].split("/")[-1]
        tone = row["Tone (Ground Truth/Prediction)"].split("/")[-1]
        cites = row["Playbook Citations"].replace("Page ", "P")
        rows.append([row["Shipment"], decision, esc, tone, row["Task Completion"], cites])
table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.45), Inches(1.1), Inches(9.1), Inches(4.9)).table
widths = [1.0, 1.85, 0.85, 0.9, 0.9, 3.6]
for i, width in enumerate(widths):
    table.columns[i].width = Inches(width)
for c, header in enumerate(headers):
    cell = table.cell(0, c)
    cell.text = header
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY
    for p in cell.text_frame.paragraphs:
        p.font.name = "Aptos"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = WHITE
for r, row in enumerate(rows, start=1):
    for c, value in enumerate(row):
        cell = table.cell(r, c)
        cell.text = value
        cell.fill.solid()
        cell.fill.fore_color.rgb = LIGHT if r % 2 else WHITE
        for p in cell.text_frame.paragraphs:
            p.font.name = "Aptos"
            p.font.size = Pt(9)
            p.font.color.rgb = GREEN if (c == 4 and value == "PASS") else GRAY
            p.font.bold = c == 4
add_body(slide, [
    "Noise scenarios intentionally show N/A for escalation and playbook citations because they are short-circuited before context retrieval.",
    "Exception scenarios cite the retrieved playbook pages used for the trace; repeated pages reflect top-k retrieval, not repeated source documents.",
], x=0.65, y=6.15, w=8.7, h=0.55, size=9, color=GRAY, spacing=1)
add_footer(slide)

# Slide 12 - conclusions and recommendations
slide = prs.slides[11]
remove_all(slide)
add_title(slide, "Conclusions and Business Recommendations")
add_card(slide, "WHAT WORKED", "All 10 unique shipment scenarios were completed correctly. Duplicate scans and noise were filtered, policy-based escalation was enforced, and customer tone matched tier expectations.", 0.65, 1.2, 4.25, 2.0, GREEN, 12)
add_card(slide, "NEXT PRODUCTION STEPS", "Persist shipment notes and supervisor queues, add temperature-controlled locker metadata, and expand adversarial tests for malformed records and API outages.", 5.1, 1.2, 4.25, 2.0, ORANGE, 12)
add_body(slide, [
    "Keep deterministic policy rules authoritative for safety and escalation.",
    "Use the LLM for flexible language tasks: classification rationale and customer-facing communication.",
    "Monitor task completion, escalation accuracy, tool-call accuracy, coherence, and latency in production.",
    "Keep the offline fallback as a reproducible evaluation and disaster-recovery path.",
], x=0.8, y=3.8, w=8.5, h=2.0, size=16, color=GRAY, spacing=9)
add_footer(slide, "Submission package verified against attached ground truth")

# Remove the template's final page-number-only slide if it survived as an
# empty trailing page; keep the substantive 12-slide submission.
if len(prs.slides) > 12:
    slide_id = prs.slides[-1]._element
    prs.part.presentation.slide_id_lst.remove(slide_id)

prs.save(OUTPUT)
print(f"Wrote {OUTPUT}")