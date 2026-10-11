import json
import math
import os
import re
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt


ROOT = Path.cwd()
NOTEBOOK = ROOT / "Project_3_Low_Code_final.ipynb"
PRESENTATION = ROOT / "Project_3_Delivery_Exception_Resolution_Presentation.pptx"
VECTOR_DB = ROOT / ".agents" / "outputs" / "playbook_chroma"


def replace_cell(notebook, marker, new_source):
    matches = []
    for i, cell in enumerate(notebook["cells"]):
        source = "".join(cell.get("source", []))
        if marker in source:
            matches.append(i)
    if len(matches) != 1:
        raise ValueError(f"Expected one cell containing {marker!r}; found {matches}")
    cell = notebook["cells"][matches[0]]
    normalized_source = new_source.splitlines(keepends=True)
    if cell.get("source") != normalized_source:
        cell["source"] = normalized_source
        if cell["cell_type"] == "code":
            cell["outputs"] = []
            cell["execution_count"] = None
        else:
            cell.pop("outputs", None)
            cell.pop("execution_count", None)


def revise_notebook():
    notebook = json.loads(NOTEBOOK.read_text())

    for cell in notebook["cells"]:
        source = "".join(cell.get("source", []))
        if "import os\nimport json\nimport csv" in source:
            source = source.replace(
                "import os\nimport json\nimport csv\nimport re\nimport time\nimport sqlite3",
                "import os\nimport json\nimport csv\nimport re\nimport time\nimport math\nimport shutil\nimport sqlite3",
            )
            source = source.replace(
                "from langchain_core.messages import SystemMessage, HumanMessage",
                "from langchain_core.messages import SystemMessage, HumanMessage\n"
                "from langchain_core.embeddings import Embeddings\n"
                "from langchain_chroma import Chroma",
            )
            cell["source"] = source.splitlines(keepends=True)

        if "The playbook chunks will now be embedded" in source:
            cell["source"] = [
                "The playbook chunks are embedded with a deterministic local TF-IDF "
                "embedding and stored in a persistent ChromaDB vector database. "
                "Cosine-similarity retrieval returns the top-k policy chunks with "
                "source-page metadata. No external embedding download is needed.\n"
            ]
            cell.pop("outputs", None)
            cell.pop("execution_count", None)

    replace_cell(
        notebook,
        "class LocalTfidfEmbeddings",
        '''class LocalTfidfEmbeddings(Embeddings):
    """Deterministic corpus-fitted TF-IDF embeddings for local Chroma retrieval."""

    STOP_WORDS = {
        "a", "an", "and", "are", "as", "at", "be", "by", "for", "from",
        "in", "is", "it", "of", "on", "or", "that", "the", "this", "to",
        "with", "your", "when", "if", "may", "must", "should", "will",
    }

    def __init__(self, documents):
        tokenized = [self._tokens(doc.page_content) for doc in documents]
        vocabulary = sorted({term for terms in tokenized for term in terms})
        self.vocabulary = vocabulary
        self.index = {term: i for i, term in enumerate(vocabulary)}
        document_frequency = defaultdict(int)
        for terms in tokenized:
            for term in set(terms):
                document_frequency[term] += 1
        total = max(len(tokenized), 1)
        self.idf = {
            term: math.log((1 + total) / (1 + document_frequency[term])) + 1
            for term in vocabulary
        }

    @classmethod
    def _tokens(cls, text):
        return [
            token for token in re.findall(r"[a-z0-9_]+", text.lower())
            if len(token) > 1 and token not in cls.STOP_WORDS
        ]

    def _embed(self, text):
        counts = defaultdict(int)
        for term in self._tokens(text):
            if term in self.index:
                counts[term] += 1
        vector = [0.0] * len(self.vocabulary)
        for term, count in counts.items():
            vector[self.index[term]] = (1 + math.log(count)) * self.idf[term]
        norm = math.sqrt(sum(value * value for value in vector))
        return [value / norm for value in vector] if norm else vector

    def embed_documents(self, texts):
        return [self._embed(text) for text in texts]

    def embed_query(self, text):
        return self._embed(text)


# Rebuild the local persistent Chroma collection so reruns cannot reuse stale chunks.
VECTOR_DB_PATH = Path(".agents/outputs/playbook_chroma")
if VECTOR_DB_PATH.exists():
    shutil.rmtree(VECTOR_DB_PATH)
VECTOR_DB_PATH.mkdir(parents=True, exist_ok=True)

embedding_function = LocalTfidfEmbeddings(playbook_chunks)
vectorstore = Chroma.from_documents(
    documents=playbook_chunks,
    embedding=embedding_function,
    collection_name="delivery_exception_playbook",
    collection_metadata={"hnsw:space": "cosine"},
    persist_directory=str(VECTOR_DB_PATH),
)
RETRIEVE_CHUNKS_NUM = 3
playbook_retriever = vectorstore.as_retriever(
    search_type="similarity",
    search_kwargs={"k": RETRIEVE_CHUNKS_NUM},
)
_retrieval_smoke_test = playbook_retriever.invoke(
    "WEATHER_DELAY PERISHABLE more than 4 hours replacement"
)
assert _retrieval_smoke_test and all(
    "page" in document.metadata for document in _retrieval_smoke_test
)
print(
    f"ChromaDB vector store ready: {len(playbook_chunks)} chunks; "
    f"TF-IDF vectors={len(embedding_function.vocabulary)} dimensions; "
    f"top-k={RETRIEVE_CHUNKS_NUM}; page metadata verified."
)''',
    )

    replace_cell(
        notebook,
        "def compute_coherence_score(pred: dict) -> dict:",
        '''def compute_coherence_score(pred: dict) -> dict:
    """Score trace coherence with the configured judge or a transparent local checklist."""
    trace = {
        "shipment_id": pred.get("shipment_id"),
        "consolidated_event": pred.get("consolidated_event", {}),
        "customer_tier": pred.get("customer_profile", {}).get("tier"),
        "escalation_signals": pred.get("escalation_signals", {}),
        "resolution_output": pred.get("resolution_output", {}),
        "critic_resolution_output": pred.get("critic_resolution_output", {}),
        "resolution_revision_count": pred.get("resolution_revision_count", 0),
        "communication_output": pred.get("communication_output", {}),
        "critic_communication_output": pred.get("critic_communication_output", {}),
        "tool_calls_log": pred.get("tool_calls_log", []),
        "playbook_pages": [
            chunk.get("page") for chunk in pred.get("playbook_context", [])
        ],
        "trajectory_log": pred.get("trajectory_log", []),
    }

    if USE_REMOTE_LLM:
        try:
            response = eval_llm.invoke([
                SystemMessage(content=REASONING_TRAJECTORY_COHERENCE_PROMPT),
                HumanMessage(content=json.dumps(trace, indent=2)),
            ])
            content = response.content.strip()
            if content.startswith("```"):
                content = content.split("\\n", 1)[1].rsplit("```", 1)[0].strip()
            parsed = json.loads(content)
            score = int(parsed.get("score", 0))
            if score not in range(1, 6):
                raise ValueError("Judge returned a score outside the 1-5 scale")
            return {
                "score": score,
                "justification": parsed.get("justification", ""),
            }
        except Exception as exc:
            return {"score": 0, "justification": f"Remote judge failed: {exc}"}

    # Credential-free path: grade observable trace structure; do not return a
    # constant score. This is a transparent proxy, not an independent LLM judge.
    log = pred.get("trajectory_log", [])
    joined = " | ".join(log).lower()
    blocked = bool(pred.get("guardrail_triggered"))
    noise = bool(pred.get("noise_override"))
    required = ["preprocessor", "finalize"]
    if blocked:
        required.append("guardrail triggered")
    elif not noise:
        required.extend(["resolution_agent", "critic_resolution"])
        if pred.get("resolution_output", {}).get("is_exception") == "YES":
            required.extend(["communication_agent", "critic_communication"])

    presence = all(marker in joined for marker in required)
    final_index = next(
        (i for i, entry in enumerate(log) if entry.lower().startswith("finalize:")),
        None,
    )
    ordered = final_index is not None and (
        not log or all(
            next((i for i, entry in enumerate(log) if marker in entry.lower()), -1)
            < final_index
            for marker in required if marker != "finalize"
        )
    )
    terminal_output = bool(pred.get("final_actions"))
    citations_present = noise or blocked or bool(pred.get("playbook_context"))
    checks = [presence, ordered, terminal_output, citations_present]
    score = 1 + round(4 * sum(checks) / len(checks))
    return {
        "score": score,
        "justification": (
            f"Offline trace checklist: {sum(checks)}/{len(checks)} checks passed "
            "(required steps, order, final action, policy context)."
        ),
    }''',
    )

    replace_cell(
        notebook,
        "**Workflow:** Preprocessor",
        '''from IPython.display import Markdown, display
# Render the workflow graph as a local, credential-free visual outline.
display(Markdown(
    "**Workflow:** Preprocessor → Orchestrator → Resolution Agent → Resolution Critic "
    "→ Communication Agent → Communication Critic → Finalize  \\n"
    "The orchestrator routes critic revisions back to the Resolution Agent and applies "
    "deterministic escalation signals before completion."
))''',
    )

    # Clarify the fallback evaluation method in notebook narrative.
    for cell in notebook["cells"]:
        source = "".join(cell.get("source", []))
        if "We first define the prompt for reasoning trajectory coherence" in source:
            cell["source"] = [
                "When remote credentials are configured, an LLM-as-a-Judge scores "
                "reasoning trajectory coherence. In the credential-free run, a "
                "deterministic checklist scores observable trace completeness and "
                "order instead; that local proxy is not represented as an independent "
                "LLM judgment.\n"
            ]
            cell.pop("outputs", None)
            cell.pop("execution_count", None)
        if '"""Use eval_llm (gpt-4o) to score reasoning trajectory coherence."""' in source:
            source = source.replace(
                '"""Use eval_llm (gpt-4o) to score reasoning trajectory coherence."""',
                '"""Use the remote judge or offline trace checklist to score coherence."""',
            )
            cell["source"] = source.splitlines(keepends=True)

    # Older generated versions accidentally assigned code-only metadata to
    # Markdown cells; normalize those cells before nbformat validation.
    for cell in notebook["cells"]:
        if cell["cell_type"] == "markdown":
            cell.pop("outputs", None)
            cell.pop("execution_count", None)

    notebook["nbformat"] = 4
    notebook["nbformat_minor"] = 5
    for index, cell in enumerate(notebook["cells"]):
        if not cell.get("id"):
            cell["id"] = f"auto-cell-{index:04d}"
    NOTEBOOK.write_text(json.dumps(notebook, indent=1, ensure_ascii=False) + "\n")


def set_cell_text(cell, text, font_size, color, bold=False, align=PP_ALIGN.LEFT):
    cell.text = text
    cell.margin_left = Inches(0.055)
    cell.margin_right = Inches(0.055)
    cell.margin_top = Inches(0.035)
    cell.margin_bottom = Inches(0.025)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.text_frame.word_wrap = True
    cell.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)


def style_summary_card(shape, heading, body_lines):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.015)
    frame.margin_bottom = Inches(0.015)
    frame.vertical_anchor = MSO_ANCHOR.TOP

    def style_paragraph(paragraph, text, font_size, bold, color, after=0):
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(after)
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)

    style_paragraph(frame.paragraphs[0], heading, 13.2, True, "183757", 2)
    for line in body_lines:
        paragraph = frame.add_paragraph()
        style_paragraph(paragraph, line, 10.7, False, "33495D", 0)


def revise_presentation():
    prs = Presentation(PRESENTATION)
    target_height = Inches(5.625)

    # Keep the supplied 16:9 canvas and scale any 4:3-authored shapes into it.
    if prs.slide_height > target_height:
        factor = target_height / prs.slide_height
        for slide in prs.slides:
            for shape in slide.shapes:
                shape.top = int(shape.top * factor)
                shape.height = int(shape.height * factor)
    elif prs.slide_height < target_height:
        factor = target_height / prs.slide_height
        for slide in prs.slides:
            for shape in slide.shapes:
                shape.top = int(shape.top * factor)
                shape.height = int(shape.height * factor)

    # Refresh visible copy to reflect the actual indexed retrieval and honest
    # offline metric qualifications.
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace(
                        "Page-aware lexical retrieval over 24 chunks. It returns source page metadata for audit citations without requiring a model download.",
                        "ChromaDB cosine search over 24 local TF-IDF vectors returns ranked policy chunks with source-page metadata; no embedding download is required.",
                    )
                    run.text = run.text.replace(
                        "Coherence: 5.00 / 5.0",
                        "Coherence: 5.00 / 5 (offline trace checklist)",
                    )
                    run.text = run.text.replace(
                        "Coherence 5.00/5.0",
                        "Coherence 5.00/5 (offline checklist)",
                    )
            shape.text_frame.word_wrap = True
            shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    # Keep long headings clear of the supplied brand mark in the upper-right.
    for slide in prs.slides:
        for shape in slide.shapes:
            if (
                shape.has_text_frame
                and shape.top < Inches(0.7)
                and shape.width > Inches(8.0)
            ):
                shape.width = Inches(7.85)
                shape.text_frame.word_wrap = False
                shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    # The template's final "Happy Learning" layout adds a blue end-card and
    # oversized background slogan that conflict with the business conclusion.
    conclusion_slide = prs.slides[11]
    if conclusion_slide.slide_layout.name == "BLANK_1":
        for rel in list(conclusion_slide.part.rels.values()):
            if rel.reltype == RT.SLIDE_LAYOUT:
                conclusion_slide.part.drop_rel(rel.rId)
        conclusion_slide.part.relate_to(prs.slide_layouts[0].part, RT.SLIDE_LAYOUT)

    # Earlier copy updates flattened the rich formatting in these metric cards;
    # restore intentional heading/body hierarchy without affecting other cards.
    slide3_textboxes = {
        shape.name: shape for shape in prs.slides[2].shapes if shape.has_text_frame
    }
    style_summary_card(
        slide3_textboxes["TextBox 196"],
        "RESULT",
        [
            "Task completion: 100%",
            "Escalation accuracy: 100% (8/8)",
            "Tool-call accuracy: 100%",
            "Coherence: 5.00 / 5 (offline trace checklist)",
        ],
    )
    slide6_textboxes = {
        shape.name: shape for shape in prs.slides[5].shapes if shape.has_text_frame
    }
    style_summary_card(
        slide6_textboxes["TextBox 224"],
        "RETRIEVAL",
        [
            "ChromaDB cosine search over 24 local TF-IDF vectors returns ranked policy chunks with source-page metadata.",
            "No external embedding download is required.",
        ],
    )
    slide9_textboxes = {
        shape.name: shape for shape in prs.slides[8].shapes if shape.has_text_frame
    }
    style_summary_card(
        slide9_textboxes["TextBox 254"],
        "AGGREGATE METRICS",
        [
            "Task completion: 100%",
            "Escalation accuracy: 100% (8/8)",
            "Tool-call accuracy: 100%",
            "Coherence: 5.00 / 5 (offline checklist)",
            "Offline latency: 0.01s (not a production benchmark)",
        ],
    )

    # Update the cover date to the current workspace date.
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace("20 September 2026", "25 September 2026")

    # Rebuild the appendix table with per-case metrics, observations, and citations.
    slide = prs.slides[10]
    old_table = next(shape for shape in slide.shapes if shape.has_table)
    old_table._element.getparent().remove(old_table._element)
    table_shape = slide.shapes.add_table(
        11, 8, Inches(0.35), Inches(0.86), Inches(9.3), Inches(4.16)
    )
    table = table_shape.table
    widths = [0.78, 1.43, 0.70, 0.55, 0.56, 0.84, 3.23, 1.21]
    for col, width in enumerate(widths):
        table.columns[col].width = Inches(width)
    rows = [
        ["Shipment", "Outcome / tone", "Esc.", "Task", "Tools", "Coherence", "Trace / observation", "Playbook"],
        ["SHP-001", "NO EXCEPTION / —", "N/A", "PASS", "PASS", "5/5", "Preprocessor → delivered-noise gate → finalize.", "N/A"],
        ["SHP-002", "RESCHEDULE / formal", "PASS", "PASS", "PASS", "5/5", "Dedup → profile/RAG → reschedule → critics → escalate.", "P2, 5, 6"],
        ["SHP-003", "RESCHEDULE / casual", "PASS", "PASS", "PASS", "5/5", "Profile/RAG → address retry → both critics accept.", "P3, 2, 3"],
        ["SHP-004", "REPLACE / formal", "PASS", "PASS", "PASS", "5/5", "Profile/RAG → replace → critics → escalate.", "P4, 4, 7"],
        ["SHP-005", "LOCKER REROUTE / casual", "PASS", "PASS", "PASS", "5/5", "Locker/RAG → reroute → critics → escalate.", "P2, 2, 3"],
        ["SHP-006", "RETURN / casual", "PASS", "PASS", "PASS", "5/5", "Profile/RAG → return to sender → critics accept.", "P5, 5, 4"],
        ["SHP-007", "REPLACE / formal", "PASS", "PASS", "PASS", "5/5", "Weather/perishable RAG → replace → escalate.", "P6, 2"],
        ["SHP-008", "RESCHEDULE / casual", "PASS", "PASS", "PASS", "5/5", "History rule → reschedule → discretionary escalation.", "P2, 5, 5"],
        ["SHP-009", "REPLACE / formal", "PASS", "PASS", "PASS", "5/5", "Damage/perishable RAG → replace → escalate.", "P4, 6, 2"],
        ["SHP-010", "NO EXCEPTION / —", "N/A", "PASS", "PASS", "5/5", "Preprocessor → routine-scan gate → finalize.", "N/A"],
    ]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = RGBColor(24, 55, 87)
                color, bold = "FFFFFF", True
                size = 8.2
            else:
                cell.fill.fore_color.rgb = (
                    RGBColor(242, 247, 250) if r % 2 else RGBColor(255, 255, 255)
                )
                color, bold = "183757", False
                size = 7.5 if c in (1, 6, 7) else 7.8
            alignment = PP_ALIGN.CENTER if c in (0, 2, 3, 4, 5) else PP_ALIGN.LEFT
            set_cell_text(cell, value, size, color, bold=bold, align=alignment)
    table.rows[0].height = Inches(0.40)
    for row_index in range(1, len(table.rows)):
        table.rows[row_index].height = Inches(0.376)

    # Remove the template footnote: metric qualifications are already stated
    # on the aggregate-metrics slide, and the table keeps noise-case N/A visible.
    for note in list(slide.shapes):
        if (
            note.has_text_frame
            and ("Noise scenarios" in note.text or "Escalation is N/A" in note.text)
        ):
            note._element.getparent().remove(note._element)

    # Keep the supplied 16:9 dimensions.
    prs.slide_height = target_height
    # The original template footer was positioned for a taller 4:3 page. Move
    # its master/layout text into the visible 16:9 footer band and keep it to
    # one line so it cannot spill past the exported PDF edge.
    for master in prs.slide_masters:
        for container in [master, *list(master.slide_layouts)]:
            for shape in container.shapes:
                if not shape.has_text_frame:
                    continue
                label = shape.text.strip()
                if label.startswith("Proprietary content."):
                    shape.left = Inches(1.75)
                    shape.top = Inches(5.46)
                    shape.width = Inches(6.5)
                    shape.height = Inches(0.13)
                    shape.text_frame.word_wrap = False
                    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                    shape.text_frame.margin_left = Inches(0)
                    shape.text_frame.margin_right = Inches(0)
                    shape.text_frame.margin_top = Inches(0)
                    shape.text_frame.margin_bottom = Inches(0)
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(6.0)
                elif label == "‹#›":
                    shape.top = Inches(5.42)
                    shape.height = Inches(0.14)
                    shape.text_frame.margin_top = Inches(0)
                    shape.text_frame.margin_bottom = Inches(0)
    prs.save(PRESENTATION)


if __name__ == "__main__":
    revise_notebook()
    revise_presentation()
    print(f"Revised notebook: {NOTEBOOK}")
    print(f"Revised presentation: {PRESENTATION}")